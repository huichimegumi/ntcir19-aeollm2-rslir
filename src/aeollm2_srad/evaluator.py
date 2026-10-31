"""Final SRAD scoring and sequential-reference pipeline (no legacy runners)."""

from __future__ import annotations

import ast
import hashlib
import json
import math
import os
import random
import re
import statistics
import sys
import time
from pathlib import Path
from typing import Dict, List
import docx
from openai import OpenAI
from . import sections as section_splitter
from .prompting import (
    BEGIN_MARKER,
    END_MARKER,
    _system_json_message,
    construct_fullcontext_readability_prompt,
    construct_global_readability_prompt,
    construct_paragraph_readability_prompt,
    construct_sequential_global_readability_prompt,
)

RUBRICS_DIR = Path("data/rubrics")
BLIND_DOCS_DIR = Path("data/reports")
OUTPUTS_DIR = Path("outputs/srad")
DEBUG_RAW_DIR = OUTPUTS_DIR / "debug_raw"
RUN_FINGERPRINT = ""
READABILITY_PARAGRAPH_ALPHA = 0.3
READABILITY_GLOBAL_ALPHA = 0.7

READABILITY_LAST_WEIGHT = 0.4

READABILITY_PARAGRAPH_MAX_TOKENS = 4096

READABILITY_GLOBAL_MAX_TOKENS = 2048

READABILITY_MISSING_RETRY_ATTEMPTS = 2

READABILITY_MAX_MISSING_RATIO = 0.10

SEQUENTIAL_GLOBAL_PROMPT_VARIANT = "previous_fulltext_and_score_free_calibration_v1"

DIM_KEY_ALIASES = {
    "comprehensiveness": ["comprehensiveness", "全面性", "comprehensive"],
    "insight": ["insight", "洞察", "洞察深度"],
    "instruction_following": ["instruction_following", "指令遵循", "instruction following"],
    "readability": ["readability", "可读性"],
}

REFERENCE_SECTION_PATTERNS = [
    r"^\s*文献来源\s*$",
    r"^\s*参考文献\s*$",
    r"^\s*参考资料\s*$",
    r"^\s*资料来源\s*$",
    r"^\s*来源\s*$",
    r"^\s*引用来源\s*$",
    r"^\s*以下来源未被直接引用\s*$",
    r"^\s*(references|bibliography|sources|works cited)\s*$",
]

def _sanitize_model_name(name: str) -> str:
    return re.sub(r"[^0-9A-Za-z._-]", "_", name)



def _append_run_log(log_path: Path, message: str) -> None:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    ts = time.strftime("%Y-%m-%d %H:%M:%S")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(f"[{ts}] {message}\n")

def _file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def _build_extra_body(model: str, thinking_mode: str) -> dict:
    """Build provider-specific thinking controls for compatible models."""
    mode = (thinking_mode or "auto").lower()
    model_lower = str(model).lower()

    if mode == "default":
        return {}

    if "qwen" in model_lower:
        if mode in {"auto", "off"}:
            return {"chat_template_kwargs": {"enable_thinking": False}}
        if mode == "on":
            return {"chat_template_kwargs": {"enable_thinking": True}}

    if "deepseek" in model_lower:
        if mode == "off":
            return {"thinking": {"type": "disabled"}}
        if mode == "on":
            return {"thinking": {"type": "enabled"}}

    return {}

def _discover_reports_from_train() -> list[int]:
    """Parse and obtain the available report/task ID list."""
    ids: list[int] = []
    if not BLIND_DOCS_DIR.exists():
        return ids
    for child in BLIND_DOCS_DIR.iterdir():
        if not child.is_dir():
            continue
        m = re.match(r"[Rr]eport(\d+)$", child.name)
        if not m:
            continue
        try:
            ids.append(int(m.group(1)))
        except ValueError:
            continue
    return sorted(set(ids))

def _normalize_doc_id(doc_id: str) -> str:
    """Normalize Doc ids such as 4, 004, Doc_4, or Doc_004 to Doc_004."""
    stem = Path(str(doc_id).strip()).stem
    m = re.fullmatch(r"(?:[Dd]oc_?)?(\d+)", stem)
    if m:
        return f"Doc_{int(m.group(1)):03d}"
    return stem

def _filter_task_files(files: list[dict], doc_filter: set[str] | None, task_id: int) -> list[dict]:
    if not doc_filter:
        return files
    selected = [fi for fi in files if _normalize_doc_id(Path(fi["path"]).stem) in doc_filter]
    found = {_normalize_doc_id(Path(fi["path"]).stem) for fi in selected}
    missing = sorted(doc_filter - found)
    if missing:
        print(f"  [Warn] Report {task_id}: requested docs not found here: {', '.join(missing)}")
    return selected

def _append_jsonl(file_path: Path, payload: dict) -> None:
    file_path.parent.mkdir(parents=True, exist_ok=True)
    with open(file_path, "a", encoding="utf-8") as f:
        f.write(json.dumps(payload, ensure_ascii=False) + "\n")

def _is_valid_number(v) -> bool:
    try:
        float(v)
        return True
    except Exception:
        return False

def read_docx_text(file_path: Path) -> str:
    document = docx.Document(str(file_path))
    return "\n".join(p.text for p in document.paragraphs)

def read_docx_paragraphs(file_path: Path) -> list[str]:
    document = docx.Document(str(file_path))
    return [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]

def readability_position_weights(n: int) -> list[float]:
    if n <= 0:
        return []
    if n == 1:
        return [1.0]
    span = 1.0 - READABILITY_LAST_WEIGHT
    return [round(1.0 - span * (i / (n - 1)), 6) for i in range(n)]

def weighted_average(scores: list[float], weights: list[float]) -> float:
    pairs = [(float(s), float(w)) for s, w in zip(scores, weights)]
    if not pairs:
        return 0.0
    weight_sum = sum(w for _, w in pairs)
    if weight_sum <= 0:
        return round(sum(s for s, _ in pairs) / len(pairs), 4)
    return round(sum(s * w for s, w in pairs) / weight_sum, 4)

def _looks_like_reference_heading(block: str) -> bool:
    first_lines = [ln.strip() for ln in str(block).splitlines() if ln.strip()][:3]
    for line in first_lines:
        normalized = re.sub(r"^[#*\-\s\d.、:：]+", "", line).strip()
        if any(re.search(pattern, normalized, flags=re.I) for pattern in REFERENCE_SECTION_PATTERNS):
            return True
    return False

def _looks_like_url_reference_block(block: str) -> bool:
    text = str(block).strip()
    if not text:
        return False
    urls = re.findall(r"https?://\S+|www\.\S+", text, flags=re.I)
    if len(urls) < 2:
        return False
    url_chars = sum(len(u) for u in urls)
    return (url_chars / max(len(text), 1)) >= 0.45

def _weighted_dim_score(criteria_items: List[Dict], scored_items: List[Dict]) -> float:
    total_w = sum(float(it.get("weight", 0)) for it in criteria_items)
    if total_w <= 0:
        vals = []
        for s in scored_items:
            try:
                vals.append(float(s.get("score", 0)))
            except Exception:
                vals.append(0.0)
        return round(sum(vals) / max(len(vals), 1), 4)

    agg = 0.0
    for c, s in zip(criteria_items, scored_items):
        try:
            sv = float(s.get("score", 0))
        except Exception:
            sv = 0.0
        agg += sv * float(c.get("weight", 0))
    return round(agg / total_w, 4)

def _find_dim_value(payload: dict, target_dim: str):
    aliases = DIM_KEY_ALIASES.get(target_dim, [target_dim])
    for alias in aliases:
        if alias in payload:
            return payload.get(alias)
    return []

def _extract_scores_fallback(text: str, criterions: Dict[str, List[Dict]]) -> dict | None:
    """Best-effort extraction of per-dimension scores from malformed/non-JSON text."""
    if not isinstance(text, str) or not text.strip():
        return None

    canonical_keys = ["comprehensiveness", "insight", "instruction_following", "readability"]
    parsed: dict = {}
    hit_any = False

    for dim in canonical_keys:
        need = len(criterions.get(dim, []))
        if need <= 0:
            parsed[dim] = []
            continue

        aliases = DIM_KEY_ALIASES.get(dim, [dim])
        segment = ""
        for alias in aliases:
            p = re.compile(
                r'"?' + re.escape(alias) + r'"?\s*:\s*(.*?)(?=(?:"?(?:comprehensiveness|insight|instruction_following|readability)"?\s*:)|$)',
                flags=re.S | re.I,
            )
            m = p.search(text)
            if m:
                segment = m.group(1)
                break

        if not segment:
            parsed[dim] = []
            continue

        nums = re.findall(r"-?\d+(?:\.\d+)?", segment)
        scores = []
        for token in nums:
            try:
                val = float(token)
            except Exception:
                continue
            val = max(0.0, min(10.0, val))
            scores.append({"score": val})
            if len(scores) >= need:
                break

        if scores:
            hit_any = True
        parsed[dim] = scores

    return parsed if hit_any else None

def _parse_json_payload(raw: str):
    payload_str = extract_json_between_markers(raw)
    if not payload_str:
        payload_str = _clean_json(raw)
    if not payload_str or not str(payload_str).strip():
        raise ValueError("No JSON found")

    try:
        return json.loads(payload_str)
    except Exception:
        try:
            cleaned_payload = re.sub(r',\s*([}\]])', r'\1', payload_str)
            return json.loads(cleaned_payload)
        except Exception:
            return ast.literal_eval(payload_str)

def _coerce_score(value) -> float | None:
    try:
        score = float(value)
    except Exception:
        return None
    return max(0.0, min(10.0, score))

def _parse_paragraph_scores(raw: str, expected_indices: list[int]) -> dict[int, float]:
    parsed = _parse_json_payload(raw)
    scores: dict[int, float] = {}

    if isinstance(parsed, list):
        for offset, item in enumerate(parsed):
            idx = expected_indices[offset] if offset < len(expected_indices) else offset + 1
            if isinstance(item, dict):
                idx = int(item.get("index", idx))
                score = _coerce_score(item.get("score"))
            else:
                score = _coerce_score(item)
            if score is not None:
                scores[idx] = score
        return scores

    if not isinstance(parsed, dict):
        raise ValueError(f"Paragraph readability payload is not a dict/list: {type(parsed)}")

    items = (
        parsed.get("paragraph_scores")
        or parsed.get("scores")
        or parsed.get("readability")
        or parsed.get("paragraphs")
        or []
    )

    if isinstance(items, dict):
        for key, value in items.items():
            m = re.search(r"\d+", str(key))
            if not m:
                continue
            score = _coerce_score(value.get("score") if isinstance(value, dict) else value)
            if score is not None:
                scores[int(m.group(0))] = score
        return scores

    if isinstance(items, list):
        for offset, item in enumerate(items):
            idx = expected_indices[offset] if offset < len(expected_indices) else offset + 1
            if isinstance(item, dict):
                idx = int(item.get("index", idx))
                score = _coerce_score(item.get("score"))
            else:
                score = _coerce_score(item)
            if score is not None:
                scores[idx] = score
        return scores

    raise ValueError("Paragraph readability payload missing scores")

def _clean_json(text: str) -> str:
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL).strip()
    m = re.search(r"```json\s*(.*?)```", text, re.DOTALL)
    if m:
        return m.group(1).strip()
    m = re.search(r"```\s*(.*?)```", text, re.DOTALL)
    if m:
        return m.group(1).strip()
    m = re.search(r"\{.*\}", text, re.DOTALL)
    if m:
        return m.group(0).strip()
    return text

def extract_json_between_markers(text: str, begin_marker: str = BEGIN_MARKER, end_marker: str = END_MARKER) -> str | None:
    """
    Extract the raw JSON string between explicit begin/end markers. Returns None if not found.
    """
    try:
        pattern = re.escape(begin_marker) + r"(.*?)" + re.escape(end_marker)
        m = re.search(pattern, text, flags=re.S)
        if m:
            return m.group(1).strip()
    except Exception:
        pass

    # Fallback: try to find the first {...} block
    m2 = re.search(r"(\{.*\})", text, flags=re.S)
    if m2:
        return m2.group(1).strip()
    return None

class ContextLengthError(RuntimeError):
    """Raised when the provider rejects a request because its context is too long."""

def _is_context_length_error(error: Exception) -> bool:
    message = str(error).lower()
    patterns = (
        "context length",
        "context_length",
        "maximum context",
        "max context",
        "context window",
        "too many tokens",
        "input is too long",
        "prompt is too long",
        "request too large",
    )
    return any(pattern in message for pattern in patterns)

def _response_usage_dict(response) -> dict:
    usage = getattr(response, "usage", None)
    if usage is None:
        return {}
    result = {}
    for key in ("prompt_tokens", "completion_tokens", "total_tokens"):
        value = getattr(usage, key, None)
        if value is not None:
            result[key] = int(value)
    return result

def _call_chat_for_content(
    client,
    *,
    model: str,
    thinking_mode: str,
    prompt: str,
    max_tokens: int,
) -> tuple[str, dict, str | None]:
    extra_body = _build_extra_body(model, thinking_mode)
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": _system_json_message()},
            {"role": "user", "content": prompt},
        ],
        temperature=0.0,
        max_tokens=max_tokens,
        extra_body=extra_body if extra_body else None,
    )
    choice = response.choices[0]
    finish_reason = getattr(choice, "finish_reason", None)
    raw = choice.message.content
    if finish_reason == "length":
        raise ContextLengthError("Provider returned finish_reason=length")
    if not raw:
        raise ValueError("Model returned empty body (None or '')")
    return raw, _response_usage_dict(response), finish_reason

def _call_with_retries(callable_fn, retry_attempts: int, label: str):
    total_attempts = max(1, int(retry_attempts) + 1)
    last_error = None
    for attempt in range(1, total_attempts + 1):
        try:
            result = callable_fn()
            return result, attempt
        except ContextLengthError:
            raise
        except Exception as error:
            if _is_context_length_error(error):
                raise ContextLengthError(str(error)) from error
            last_error = error
            if attempt >= total_attempts:
                break
            wait_seconds = min(2 ** (attempt - 1), 4)
            print(
                f"[Warn] {label} failed on attempt {attempt}/{total_attempts}: "
                f"{error}; retrying in {wait_seconds}s"
            )
            time.sleep(wait_seconds)
    raise RuntimeError(f"{label} failed after {total_attempts} attempts: {last_error}") from last_error

def _parse_single_dimension_payload_for_sequential(
    raw: str,
    *,
    dim: str,
    model: str,
    prompt: str,
    doc_id: str,
    criterions: dict,
) -> dict:
    payload_str = extract_json_between_markers(raw) or _clean_json(raw)
    parsed = None
    if payload_str:
        try:
            parsed = json.loads(payload_str)
        except Exception:
            try:
                parsed = json.loads(re.sub(r",\s*([}\]])", r"\1", payload_str))
            except Exception:
                try:
                    parsed = ast.literal_eval(payload_str)
                except Exception:
                    parsed = _extract_scores_fallback(payload_str, criterions)
                    if parsed is None:
                        parsed = _extract_scores_fallback(raw, criterions)

    if parsed is None:
        DEBUG_RAW_DIR.mkdir(parents=True, exist_ok=True)
        ts = int(time.time() * 1000)
        debug_file = (
            DEBUG_RAW_DIR
            / f"{_sanitize_model_name(model)}_{_sanitize_model_name(doc_id)}_{dim}_{ts}.txt"
        )
        with open(debug_file, "w", encoding="utf-8") as stream:
            stream.write(raw if isinstance(raw, str) else str(raw))
            stream.write(f"\n\nPROMPT_LENGTH: {len(prompt)}")
        raise ValueError(f"Failed to parse {dim} output. Saved raw to {debug_file}")

    if isinstance(parsed, list):
        return {dim: parsed}
    if isinstance(parsed, (int, float)):
        return {dim: [parsed]}
    if not isinstance(parsed, dict):
        raise ValueError(f"Parsed {dim} payload is not a dict, got {type(parsed)}")

    payload = parsed
    if "judge_scores" in payload and isinstance(payload.get("judge_scores"), dict):
        judge_scores = payload["judge_scores"]
        if model in judge_scores and isinstance(judge_scores.get(model), dict):
            payload = judge_scores[model]
        elif dim in judge_scores:
            payload = judge_scores
        else:
            first_dict = next(
                (value for value in judge_scores.values() if isinstance(value, dict)),
                None,
            )
            if first_dict is not None:
                payload = first_dict
    elif "scores" in payload and isinstance(payload.get("scores"), dict):
        payload = payload["scores"]
    elif payload.get("dimension") == dim and isinstance(payload.get("scores"), list):
        payload = {dim: payload["scores"]}
    elif isinstance(payload.get("scores"), list):
        payload = {dim: payload["scores"]}
    return payload

def _global_score_from_payload(payload: dict, rubric: dict) -> tuple[float, list[dict]]:
    criteria_items = rubric.get("criterions", {}).get("readability", [])
    dim_val = _find_dim_value(payload, "readability")
    scored_list = dim_val if isinstance(dim_val, list) else []
    normalized = []
    for idx, criterion in enumerate(criteria_items, 1):
        source = scored_list[idx - 1] if idx - 1 < len(scored_list) else {}
        if isinstance(source, dict):
            raw_score = source.get("score", 0.0)
            raw_reason = str(source.get("reason", ""))
        else:
            raw_score = source
            raw_reason = ""
        score = _coerce_score(raw_score)
        if score is None:
            score = 0.0
        normalized.append(
            {
                "criterionId": f"readability_global_{idx}",
                "criterion": criterion.get("criterion", ""),
                "weight": float(criterion.get("weight", 0)),
                "score": round(score, 4),
                "reason": raw_reason,
            }
        )
    return _weighted_dim_score(criteria_items, normalized), normalized

def evaluate_section_readability_once(
    doc_path: Path,
    rubric: dict,
    client,
    model: str,
    thinking_mode: str,
) -> tuple[float, dict]:
    """Run the original section-readability method once, without Global scoring."""
    doc_text = read_docx_text(doc_path)
    if not doc_text.strip():
        raise ValueError(f"Document is empty: {doc_path}")

    paragraphs = read_docx_paragraphs(doc_path)
    section_dicts, split_strategy, dropped_reference_paragraphs = (
        section_splitter.build_readability_sections(sys.modules[__name__], paragraphs)
    )
    readability_blocks = [section["text"] for section in section_dicts]
    if not readability_blocks:
        section_dicts = [
            {
                "title": "Full Document",
                "start_para": 1,
                "end_para": len(paragraphs),
                "text": doc_text.strip(),
            }
        ]
        readability_blocks = [doc_text.strip()]

    indexed_blocks = list(enumerate(readability_blocks, 1))
    block_scores_by_index: dict[int, float] = {}

    def score_batch(batch: list[tuple[int, str]], attempt_label: str) -> dict[int, float]:
        expected_indices = [idx for idx, _ in batch]
        if attempt_label == "fullcontext":
            prompt = construct_fullcontext_readability_prompt(doc_text, rubric, batch)
        else:
            prompt = construct_paragraph_readability_prompt(
                rubric,
                batch,
                len(readability_blocks),
            )
        print(
            f"[Debug] Calling {model} for {doc_path.name} / readability sections "
            f"{expected_indices[0]}-{expected_indices[-1]} [{attempt_label}] "
            f"(Length: {len(prompt)})"
        )
        raw, _, _ = _call_chat_for_content(
            client,
            model=model,
            thinking_mode=thinking_mode,
            prompt=prompt,
            max_tokens=READABILITY_PARAGRAPH_MAX_TOKENS,
        )
        parsed_scores = _parse_paragraph_scores(raw, expected_indices)
        expected_set = set(expected_indices)
        return {
            idx: score
            for idx, score in parsed_scores.items()
            if idx in expected_set
        }

    block_scores_by_index.update(score_batch(indexed_blocks, "fullcontext"))
    missing_after_fullcontext = [
        idx for idx, _ in indexed_blocks if idx not in block_scores_by_index
    ]
    retry_attempt = 1
    while (
        missing_after_fullcontext
        and retry_attempt <= READABILITY_MISSING_RETRY_ATTEMPTS
    ):
        retry_batch = [
            (idx, readability_blocks[idx - 1])
            for idx in missing_after_fullcontext
        ]
        print(
            f"[Debug] Retrying {model} for {doc_path.name} / missing sections "
            f"{missing_after_fullcontext} (attempt {retry_attempt})"
        )
        block_scores_by_index.update(
            score_batch(retry_batch, f"retry-{retry_attempt}")
        )
        missing_after_fullcontext = [
            idx for idx, _ in indexed_blocks if idx not in block_scores_by_index
        ]
        retry_attempt += 1

    missing = [idx for idx, _ in indexed_blocks if idx not in block_scores_by_index]
    if missing:
        missing_ratio = len(missing) / max(len(indexed_blocks), 1)
        if missing_ratio > READABILITY_MAX_MISSING_RATIO:
            raise ValueError(
                f"Missing readability section scores for sections {missing}; "
                f"missing_ratio={missing_ratio:.4f}"
            )

    all_weights = readability_position_weights(len(indexed_blocks))
    scored_indices = [
        idx for idx, _ in indexed_blocks if idx in block_scores_by_index
    ]
    scores = [block_scores_by_index[idx] for idx in scored_indices]
    weights = [all_weights[idx - 1] for idx in scored_indices]
    section_score = weighted_average(scores, weights)
    details = {
        "section_blocks": [
            {
                "index": idx,
                "title": section_dicts[idx - 1].get("title", ""),
                "start_para": section_dicts[idx - 1].get("start_para", ""),
                "end_para": section_dicts[idx - 1].get("end_para", ""),
                "weight": all_weights[idx - 1],
                "score": round(block_scores_by_index[idx], 4),
            }
            for idx, _ in indexed_blocks
            if idx in block_scores_by_index
        ],
        "split_strategy": split_strategy,
        "dropped_reference_paragraphs": dropped_reference_paragraphs,
        "dropped_missing_blocks": missing,
        "section_readability": section_score,
    }
    return section_score, details

def evaluate_global_readability_prompt(
    *,
    prompt: str,
    doc_id: str,
    rubric: dict,
    client,
    model: str,
    thinking_mode: str,
) -> tuple[float, dict]:
    raw, usage, finish_reason = _call_chat_for_content(
        client,
        model=model,
        thinking_mode=thinking_mode,
        prompt=prompt,
        max_tokens=READABILITY_GLOBAL_MAX_TOKENS,
    )
    payload = _parse_single_dimension_payload_for_sequential(
        raw,
        dim="readability",
        model=model,
        prompt=prompt,
        doc_id=doc_id,
        criterions=rubric.get("criterions", {}),
    )
    score, criteria = _global_score_from_payload(payload, rubric)
    return score, {
        "criteria": criteria,
        "prompt_chars": len(prompt),
        "usage": usage,
        "finish_reason": finish_reason,
    }

def _load_jsonl_latest(path: Path, key_fn) -> dict:
    latest = {}
    if not path.exists():
        return latest
    with open(path, "r", encoding="utf-8") as stream:
        for line_number, line in enumerate(stream, 1):
            if not line.strip():
                continue
            try:
                payload = json.loads(line)
                latest[key_fn(payload)] = payload
            except Exception as error:
                print(f"[Warn] Ignore malformed JSONL line {path}:{line_number}: {error}")
    return latest

def _write_json_atomic(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(path.name + ".tmp")
    with open(temp_path, "w", encoding="utf-8") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2)
    os.replace(temp_path, path)

def _upsert_readability_baseline(
    path: Path,
    rows: list[tuple[str, str, float]],
) -> None:
    header = (
        "taskId\tquestionId\tanswerId\tcomprehensiveness_score\tinsight_score\t"
        "instruction_following_score\treadability_score\n"
    )
    existing: dict[tuple[str, str], str] = {}
    if path.exists():
        with open(path, "r", encoding="utf-8") as stream:
            for line in stream:
                if not line.strip() or line.startswith("taskId\t"):
                    continue
                parts = line.rstrip("\n").split("\t")
                if len(parts) >= 7:
                    existing[(parts[1], parts[2])] = line.rstrip("\n")

    for question_id, doc_id, score in rows:
        existing[(str(question_id), str(doc_id))] = (
            f"0\t{question_id}\t{doc_id}\t\t\t\t{round(float(score), 4)}"
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as stream:
        stream.write(header)
        def sort_key(item):
            question_id, doc_id = item
            question_part = (
                (0, int(question_id))
                if str(question_id).isdigit()
                else (1, str(question_id))
            )
            return question_part, _normalize_doc_id(doc_id)

        for key in sorted(existing, key=sort_key):
            stream.write(existing[key] + "\n")

def _processed_readability_keys(path: Path) -> set[tuple[str, str]]:
    keys = set()
    if not path.exists():
        return keys
    with open(path, "r", encoding="utf-8") as stream:
        for line in stream:
            if not line.strip() or line.startswith("taskId\t"):
                continue
            parts = line.rstrip("\n").split("\t")
            if len(parts) >= 7 and _is_valid_number(parts[6]):
                keys.add((str(parts[1]), str(parts[2])))
    return keys

def _method_signature(
    *,
    model: str,
    thinking_mode: str,
    global_order_seeds: list[int],
) -> str:
    script_hash = _file_sha256(Path(__file__).resolve())[:16]
    return "|".join(
        [
            SEQUENTIAL_GLOBAL_PROMPT_VARIANT,
            script_hash,
            RUN_FINGERPRINT,
            str(model),
            str(thinking_mode),
            f"last_weight={READABILITY_LAST_WEIGHT}",
            f"section_alpha={READABILITY_PARAGRAPH_ALPHA}",
            f"order_seeds={','.join(str(seed) for seed in global_order_seeds)}",
        ]
    )

def _document_sort_key(file_info: dict):
    doc_id = _normalize_doc_id(Path(file_info["path"]).stem)
    match = re.search(r"(\d+)$", doc_id)
    return int(match.group(1)) if match else doc_id

def _select_report_files(
    task_id: int,
    *,
    doc_filter: set[str] | None,
    sample_size: int,
    sample_seed: int,
) -> list[dict]:
    files = [
        file_info
        for file_info in get_task_files(task_id)
        if Path(file_info["path"]).exists()
    ]
    files = _filter_task_files(files, doc_filter, task_id)
    files = sorted(files, key=_document_sort_key)
    if sample_size > 0 and len(files) > sample_size:
        rng = random.Random(int(sample_seed) * 100000 + int(task_id))
        files = sorted(rng.sample(files, sample_size), key=_document_sort_key)
    return files

def _build_or_load_chain_orders(
    orders_file: Path,
    *,
    method_signature: str,
    task_id: int,
    doc_ids: list[str],
    global_order_seeds: list[int],
) -> tuple[dict, dict[int, list[str]]]:
    payload = {
        "method_signature": method_signature,
        "global_order_seeds": list(global_order_seeds),
        "reports": {},
    }
    if orders_file.exists():
        try:
            with open(orders_file, "r", encoding="utf-8") as stream:
                existing = json.load(stream)
            if existing.get("method_signature") == method_signature:
                payload = existing
        except Exception as error:
            print(f"[Warn] Rebuilding malformed chain order file: {error}")

    report_key = str(task_id)
    saved_report = payload.setdefault("reports", {}).setdefault(report_key, {})
    orders: dict[int, list[str]] = {}
    expected_set = set(doc_ids)
    used_first_docs: set[str] = set()
    for seed in global_order_seeds:
        seed_key = str(seed)
        saved_order = saved_report.get(seed_key)
        if (
            isinstance(saved_order, list)
            and len(saved_order) == len(doc_ids)
            and set(saved_order) == expected_set
        ):
            order = [str(doc_id) for doc_id in saved_order]
        else:
            order = list(doc_ids)
            rng = random.Random(int(seed) * 100000 + int(task_id))
            rng.shuffle(order)
            if order and order[0] in used_first_docs and len(used_first_docs) < len(order):
                replacement_index = next(
                    idx
                    for idx, doc_id in enumerate(order[1:], 1)
                    if doc_id not in used_first_docs
                )
                order[0], order[replacement_index] = (
                    order[replacement_index],
                    order[0],
                )
            saved_report[seed_key] = order
        orders[int(seed)] = order
        if order:
            used_first_docs.add(order[0])

    _write_json_atomic(orders_file, payload)
    return payload, orders

def run_sequential_global_evaluation(
    tasks: list[int],
    api_config: dict,
    *,
    model: str,
    workers: int,
    doc_filter: set[str] | None,
    thinking_mode: str,
    run_id: str | None,
    repeat_count: int,
    global_order_seeds: list[int],
    sample_size: int,
    sample_seed: int,
    api_retry_attempts: int,
    client=None,
) -> None:
    if client is None:
        client = OpenAI(api_key=api_config["api_key"], base_url=api_config["api_base"])

    experiment_dir = OUTPUTS_DIR
    combined_file = experiment_dir / "predictions" / "srad.tsv"
    global_only_file = experiment_dir / "diagnostics" / "global.tsv"
    section_only_file = experiment_dir / "diagnostics" / "section.tsv"
    run_log_file = experiment_dir / "run.log"
    params_file = experiment_dir / "experiment_params.json"
    section_state_file = experiment_dir / "section_repeat_scores.jsonl"
    global_chain_file = experiment_dir / "global_chain_scores.jsonl"
    document_summary_file = experiment_dir / "document_summary.jsonl"
    chain_orders_file = experiment_dir / "chain_orders.json"

    signature = _method_signature(
        model=model,
        thinking_mode=thinking_mode,
        global_order_seeds=global_order_seeds,
    )
    previous_method_signature = None
    if params_file.exists():
        try:
            with open(params_file, "r", encoding="utf-8") as stream:
                previous_params = json.load(stream)
            previous_method_signature = (
                previous_params.get("sequential_global", {}).get("method_signature")
            )
        except Exception:
            previous_method_signature = None

    _write_sequential_experiment_params(
        params_file,
        combined_file=combined_file,
        global_only_file=global_only_file,
        section_only_file=section_only_file,
        run_log_file=run_log_file,
        section_state_file=section_state_file,
        global_chain_file=global_chain_file,
        document_summary_file=document_summary_file,
        chain_orders_file=chain_orders_file,
        tasks=tasks,
        api_config=api_config,
        model=model,
        workers=workers,
        doc_filter=doc_filter,
        thinking_mode=thinking_mode,
        run_id=run_id,
        repeat_count=repeat_count,
        global_order_seeds=global_order_seeds,
        sample_size=sample_size,
        sample_seed=sample_seed,
        api_retry_attempts=api_retry_attempts,
        method_signature=signature,
    )
    _append_run_log(
        run_log_file,
        f"START sequential_global model={model} tasks={tasks} "
        f"repeat_count={repeat_count} order_seeds={global_order_seeds} "
        f"thinking_mode={thinking_mode} workers={workers}",
    )

    section_state = _load_jsonl_latest(
        section_state_file,
        lambda payload: (
            payload.get("method_signature"),
            int(payload.get("report_id", -1)),
            str(payload.get("doc_id", "")),
            int(payload.get("repeat_index", -1)),
        ),
    )
    global_state = _load_jsonl_latest(
        global_chain_file,
        lambda payload: (
            payload.get("method_signature"),
            int(payload.get("report_id", -1)),
            int(payload.get("order_seed", -1)),
            str(payload.get("doc_id", "")),
        ),
    )

    combined_processed = _processed_readability_keys(combined_file)
    global_processed = _processed_readability_keys(global_only_file)
    section_processed = _processed_readability_keys(section_only_file)

    for task_id in tasks:
        print(f"\n{'=' * 60}")
        print(f"Report {task_id}: Section repeat average + Sequential Global reference")
        print(f"{'=' * 60}")
        rubric = load_rubric(task_id)
        if not rubric:
            print(f"  [Warn] Report {task_id}: rubric missing, skipped")
            _append_run_log(run_log_file, f"REPORT_SKIP report={task_id} reason=rubric_missing")
            continue

        files = _select_report_files(
            task_id,
            doc_filter=doc_filter,
            sample_size=sample_size,
            sample_seed=sample_seed,
        )
        if not files:
            print("  本任务无待评估文档")
            _append_run_log(run_log_file, f"REPORT_SKIP report={task_id} reason=no_documents")
            continue

        question_id = str(rubric.get("id", task_id))
        doc_by_id = {
            _normalize_doc_id(Path(file_info["path"]).stem): file_info
            for file_info in files
        }
        doc_ids = sorted(doc_by_id, key=lambda value: int(re.search(r"\d+$", value).group()))
        expected_keys = {(question_id, doc_id) for doc_id in doc_ids}
        if (
            previous_method_signature == signature
            and
            expected_keys <= combined_processed
            and expected_keys <= global_processed
            and expected_keys <= section_processed
        ):
            print("  本任务三套 baseline 均已完整，断点续跑跳过")
            _append_run_log(
                run_log_file,
                f"REPORT_SKIP report={task_id} reason=all_outputs_complete",
            )
            continue

        _, orders = _build_or_load_chain_orders(
            chain_orders_file,
            method_signature=signature,
            task_id=task_id,
            doc_ids=doc_ids,
            global_order_seeds=global_order_seeds,
        )
        doc_texts = {
            doc_id: read_docx_text(Path(doc_by_id[doc_id]["path"]))
            for doc_id in doc_ids
        }

        try:
            section_scores: dict[str, float] = {}
            section_repeat_values: dict[str, list[float]] = {}
            for doc_id in doc_ids:
                doc_path = Path(doc_by_id[doc_id]["path"])
                repeat_values = []
                for repeat_index in range(1, repeat_count + 1):
                    state_key = (signature, int(task_id), doc_id, repeat_index)
                    existing = section_state.get(state_key)
                    if existing and _is_valid_number(existing.get("score")):
                        score = float(existing["score"])
                    else:
                        label = (
                            f"Report {task_id} {doc_id} section repeat "
                            f"{repeat_index}/{repeat_count}"
                        )
                        (result, attempts) = _call_with_retries(
                            lambda path=doc_path: evaluate_section_readability_once(
                                path,
                                rubric,
                                client,
                                model,
                                thinking_mode,
                            ),
                            api_retry_attempts,
                            label,
                        )
                        score, details = result
                        existing = {
                            "method_signature": signature,
                            "prompt_variant": "original_section_readability",
                            "report_id": int(task_id),
                            "doc_id": doc_id,
                            "repeat_index": repeat_index,
                            "score": round(float(score), 4),
                            "api_attempts": attempts,
                            "details": details,
                            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
                        }
                        _append_jsonl(section_state_file, existing)
                        section_state[state_key] = existing
                    repeat_values.append(float(existing["score"]))
                section_repeat_values[doc_id] = repeat_values
                section_scores[doc_id] = round(
                    sum(repeat_values) / len(repeat_values),
                    4,
                )
                _append_run_log(
                    run_log_file,
                    f"SECTION_SUCCESS report={task_id} doc={doc_id} "
                    f"score={section_scores[doc_id]} repeats={repeat_values}",
                )

            global_values: dict[str, list[float]] = {doc_id: [] for doc_id in doc_ids}
            global_entries: dict[str, list[dict]] = {doc_id: [] for doc_id in doc_ids}
            for seed in global_order_seeds:
                order = orders[int(seed)]
                previous_doc_id = None
                previous_score = None
                for position, doc_id in enumerate(order, 1):
                    state_key = (signature, int(task_id), int(seed), doc_id)
                    existing = global_state.get(state_key)
                    expected_reference_id = previous_doc_id
                    expected_reference_score = (
                        None if previous_score is None else round(float(previous_score), 4)
                    )
                    reusable = bool(
                        existing
                        and _is_valid_number(existing.get("score"))
                        and int(existing.get("position", -1)) == position
                        and existing.get("reference_doc_id") == expected_reference_id
                        and (
                            expected_reference_score is None
                            or (
                                _is_valid_number(existing.get("reference_score"))
                                and abs(
                                    float(existing["reference_score"])
                                    - expected_reference_score
                                )
                                <= 1e-4
                            )
                        )
                    )

                    if not reusable:
                        current_text = doc_texts[doc_id]
                        fallback = False
                        fallback_reason = None
                        if previous_doc_id is None:
                            prompt = construct_global_readability_prompt(
                                current_text,
                                rubric,
                            )
                            mode = "independent_initial"
                            (result, attempts) = _call_with_retries(
                                lambda p=prompt, d=doc_id: evaluate_global_readability_prompt(
                                    prompt=p,
                                    doc_id=d,
                                    rubric=rubric,
                                    client=client,
                                    model=model,
                                    thinking_mode=thinking_mode,
                                ),
                                api_retry_attempts,
                                f"Report {task_id} seed {seed} {doc_id} initial global",
                            )
                        else:
                            prompt = construct_sequential_global_readability_prompt(
                                doc_texts[previous_doc_id],
                                previous_doc_id,
                                float(previous_score),
                                current_text,
                                doc_id,
                                rubric,
                            )
                            mode = "sequential_reference"
                            try:
                                (result, attempts) = _call_with_retries(
                                    lambda p=prompt, d=doc_id: evaluate_global_readability_prompt(
                                        prompt=p,
                                        doc_id=d,
                                        rubric=rubric,
                                        client=client,
                                        model=model,
                                        thinking_mode=thinking_mode,
                                    ),
                                    api_retry_attempts,
                                    f"Report {task_id} seed {seed} {doc_id} referenced global",
                                )
                            except ContextLengthError as error:
                                fallback = True
                                fallback_reason = str(error)
                                mode = "independent_context_fallback"
                                fallback_prompt = construct_global_readability_prompt(
                                    current_text,
                                    rubric,
                                )
                                (result, attempts) = _call_with_retries(
                                    lambda p=fallback_prompt, d=doc_id: evaluate_global_readability_prompt(
                                        prompt=p,
                                        doc_id=d,
                                        rubric=rubric,
                                        client=client,
                                        model=model,
                                        thinking_mode=thinking_mode,
                                    ),
                                    api_retry_attempts,
                                    f"Report {task_id} seed {seed} {doc_id} context fallback",
                                )

                        score, details = result
                        existing = {
                            "method_signature": signature,
                            "prompt_variant": SEQUENTIAL_GLOBAL_PROMPT_VARIANT,
                            "report_id": int(task_id),
                            "order_seed": int(seed),
                            "position": position,
                            "doc_id": doc_id,
                            "reference_doc_id": previous_doc_id,
                            "reference_score": expected_reference_score,
                            "score": round(float(score), 4),
                            "mode": mode,
                            "context_fallback": fallback,
                            "fallback_reason": fallback_reason,
                            "api_attempts": attempts,
                            "details": details,
                            "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
                        }
                        _append_jsonl(global_chain_file, existing)
                        global_state[state_key] = existing
                        _append_run_log(
                            run_log_file,
                            f"GLOBAL_SUCCESS report={task_id} seed={seed} position={position} "
                            f"doc={doc_id} reference={previous_doc_id} score={score} "
                            f"mode={mode} prompt_chars={details.get('prompt_chars')}",
                        )

                    score = float(existing["score"])
                    global_values[doc_id].append(score)
                    global_entries[doc_id].append(existing)
                    previous_doc_id = doc_id
                    previous_score = score

            combined_rows = []
            global_rows = []
            section_rows = []
            for doc_id in doc_ids:
                values = global_values[doc_id]
                if len(values) != len(global_order_seeds):
                    raise ValueError(
                        f"{doc_id} has {len(values)} Global scores; "
                        f"expected {len(global_order_seeds)}"
                    )
                global_score = round(sum(values) / len(values), 4)
                global_std = round(statistics.pstdev(values), 4) if len(values) > 1 else 0.0
                section_score = section_scores[doc_id]
                combined_score = round(
                    READABILITY_PARAGRAPH_ALPHA * section_score
                    + READABILITY_GLOBAL_ALPHA * global_score,
                    4,
                )
                fallback_count = sum(
                    bool(entry.get("context_fallback"))
                    for entry in global_entries[doc_id]
                )
                summary = {
                    "method_signature": signature,
                    "report_id": int(task_id),
                    "question_id": question_id,
                    "doc_id": doc_id,
                    "section_repeat_scores": section_repeat_values[doc_id],
                    "section_score": section_score,
                    "global_order_seeds": list(global_order_seeds),
                    "global_order_scores": values,
                    "global_score": global_score,
                    "global_order_std": global_std,
                    "context_fallback_count": fallback_count,
                    "combined_score": combined_score,
                    "section_alpha": READABILITY_PARAGRAPH_ALPHA,
                    "global_alpha": READABILITY_GLOBAL_ALPHA,
                    "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
                }
                _append_jsonl(document_summary_file, summary)
                combined_rows.append((question_id, doc_id, combined_score))
                global_rows.append((question_id, doc_id, global_score))
                section_rows.append((question_id, doc_id, section_score))
                _append_run_log(
                    run_log_file,
                    f"DOC_SUCCESS report={task_id} doc={doc_id} section={section_score} "
                    f"global={global_score} global_std={global_std} "
                    f"combined={combined_score} context_fallbacks={fallback_count}",
                )

            _upsert_readability_baseline(combined_file, combined_rows)
            _upsert_readability_baseline(global_only_file, global_rows)
            _upsert_readability_baseline(section_only_file, section_rows)
            combined_processed.update(expected_keys)
            global_processed.update(expected_keys)
            section_processed.update(expected_keys)
            _append_run_log(
                run_log_file,
                f"REPORT_SUCCESS report={task_id} documents={len(doc_ids)}",
            )
        except Exception as error:
            print(f"  [Error] Report {task_id}: {error}")
            _append_run_log(
                run_log_file,
                f"REPORT_ERROR report={task_id} error={repr(error)}",
            )
            raise RuntimeError(f"SRAD failed for Report {task_id}") from error

    print("\nSequential Global reference 评估完成。")
    print(f"Combined baseline: {combined_file}")
    print(f"Global-only baseline: {global_only_file}")
    print(f"Section-only baseline: {section_only_file}")
    print(f"实验参数: {params_file}")
    print(f"运行日志: {run_log_file}")
    _append_run_log(run_log_file, "END sequential_global")


def get_task_files(task_id: int) -> list[dict]:
    """Discover official blind IDs without the private model-name mapping workbook."""
    directories = [p for p in BLIND_DOCS_DIR.iterdir()
                   if p.is_dir() and re.fullmatch(rf'[Rr]eport0*{task_id}', p.name)]
    if len(directories) != 1:
        raise ValueError(f"Expected one Report{task_id} directory")
    files, seen = [], set()
    for path in sorted(directories[0].glob('*.docx')):
        if path.name.startswith('~$'):
            continue
        if not re.fullmatch(r'Doc_\d{3,}', path.stem):
            raise ValueError(f"Expected official Doc_### ID: {path.name}")
        doc_id = _normalize_doc_id(path.stem)
        if doc_id != path.stem or doc_id in seen:
            raise ValueError(f"Noncanonical or duplicate document ID: {path.name}")
        seen.add(doc_id)
        files.append({'path': path, 'task_id': task_id})
    return files


def load_rubric(task_id: int) -> dict:
    for name in (f'criteria{task_id}.json', f'criterion{task_id}.json'):
        path = RUBRICS_DIR / name
        if not path.is_file():
            continue
        rubric = json.loads(path.read_text(encoding='utf-8-sig'))
        if int(rubric.get('id', task_id)) != task_id:
            raise ValueError(f'Rubric ID mismatch: {path.name}')
        criteria = rubric.get('criterions', {}).get('readability', [])
        if not criteria:
            raise ValueError(f'Missing readability criteria: {path.name}')
        weights = [float(c.get('weight', 0)) for c in criteria]
        if not all(math.isfinite(w) and w >= 0 for w in weights) or sum(weights) <= 0:
            raise ValueError(f'Invalid readability weights: {path.name}')
        return rubric
    raise FileNotFoundError(f'Missing rubric for Report{task_id}')


def _write_sequential_experiment_params(params_path: Path, **kwargs) -> None:
    """Record method settings without serializing credentials or legacy paths."""
    payload = {
        'sequential_global': {'method_signature': kwargs['method_signature']},
        'model': kwargs['model'], 'thinking_mode': kwargs['thinking_mode'],
        'temperature': 0, 'section_alpha': READABILITY_PARAGRAPH_ALPHA,
        'global_alpha': READABILITY_GLOBAL_ALPHA, 'last_weight': READABILITY_LAST_WEIGHT,
        'section_repeat_count': kwargs['repeat_count'],
        'global_order_seeds': kwargs['global_order_seeds'],
        'section_max_tokens': READABILITY_PARAGRAPH_MAX_TOKENS,
        'global_max_tokens': READABILITY_GLOBAL_MAX_TOKENS,
        'section_missing_retry_attempts': READABILITY_MISSING_RETRY_ATTEMPTS,
        'section_max_missing_ratio': READABILITY_MAX_MISSING_RATIO,
        'context_fallback': 'independent_global_on_provider_context_error_or_length_finish',
        'order_generation': 'seed_times_100000_plus_question_id_shuffle_distinct_first_when_possible',
        'prompt_variant': SEQUENTIAL_GLOBAL_PROMPT_VARIANT,
        'run_fingerprint': RUN_FINGERPRINT,
    }
    _write_json_atomic(params_path, payload)
