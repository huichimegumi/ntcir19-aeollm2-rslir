from docx import Document

from aeollm2_osodt.document import read_docx_with_tables


def test_reader_preserves_paragraph_table_paragraph_order(tmp_path):
    path = tmp_path / "Doc_001.docx"
    doc = Document()
    doc.add_paragraph("before")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    doc.add_paragraph("after")
    doc.save(path)

    record = read_docx_with_tables(path, question_id=1)
    assert record.text.index("before") < record.text.index("[TABLE_START]")
    assert record.text.index("[TABLE_END]") < record.text.index("after")
    assert "A\tB" in record.text
    assert "C\tD" in record.text
    assert record.table_count == 1
    assert record.table_row_count == 2
