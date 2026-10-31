"""Synthetic tests for the final SRAD protocol; no network or official data."""
import json
from pathlib import Path
from types import SimpleNamespace

import docx
import pandas as pd
import pytest

from aeollm2_srad import evaluator as ev, prompting, sections
from aeollm2_srad.cli import configure

ROOT = Path(__file__).parents[1]


@pytest.fixture
def prepared(tmp_path):
    reports, rubrics, out = tmp_path / 'reports', tmp_path / 'rubrics', tmp_path / 'outputs'
    (reports / 'Report1').mkdir(parents=True)
    rubrics.mkdir()
    rubric = {'id': 1, 'prompt': 'Compare the evidence.', 'criterions': {'readability': [
        {'criterion': 'Clarity', 'explanation': 'Clear prose', 'weight': 1}]}}
    (rubrics / 'criterion1.json').write_text(json.dumps(rubric))
    for i in range(1, 4):
        d = docx.Document()
        d.add_paragraph(f'Synthetic report {i}.')
        d.add_paragraph('A clear explanation. ' * 40)
        d.save(reports / 'Report1' / f'Doc_{i:03d}.docx')
    p = configure(ROOT / 'configs/srad-deepseek-v4-flash-paper.yaml', reports, rubrics, out)
    return p, reports, rubrics, out


class Judge:
    def __init__(self):
        self.chat = SimpleNamespace(completions=self)
        self.requests = []
        self.global_count = 0

    def create(self, **kwargs):
        self.requests.append(kwargs)
        prompt = kwargs['messages'][1]['content']
        assert kwargs['temperature'] == 0
        assert kwargs['extra_body'] == {'thinking': {'type': 'disabled'}}
        if 'paragraph_scores' in prompt:
            import re
            indices = [int(x) for x in re.findall(r'\[文章部分 (\d+)/\d+\]', prompt)]
            raw = {'paragraph_scores': [{'index': i, 'score': 6 + i / 10} for i in indices]}
            assert kwargs['max_tokens'] == 4096
        else:
            self.global_count += 1
            raw = {'readability': [4 + self.global_count / 10]}
            assert kwargs['max_tokens'] == 2048
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(
            content=json.dumps(raw)))], usage=None)


def run(p, judge):
    ev.run_sequential_global_evaluation(p['tasks'], {'api_base': 'https://fake.example', 'api_key': 'fixture'},
        model='deepseek-v4-flash', workers=1, doc_filter=None, thinking_mode='off', run_id=None,
        repeat_count=3, global_order_seeds=[42, 43, 44], sample_size=0, sample_seed=42,
        api_retry_attempts=0, client=judge)


def test_pipeline_references_fusion_and_resume(prepared):
    p, _, _, out = prepared
    judge = Judge()
    run(p, judge)
    assert len(judge.requests) == 18  # 3 docs * (3 section repeats + 3 global orders)
    chain = [json.loads(x) for x in (out / 'global_chain_scores.jsonl').read_text().splitlines()]
    assert len(chain) == 9
    assert len({x['doc_id'] for x in chain if x['position'] == 1}) == 3
    for i, entry in enumerate(chain):
        if entry['position'] == 1:
            assert entry['reference_doc_id'] is None
        else:
            assert entry['reference_doc_id'] == chain[i-1]['doc_id']
            assert entry['reference_score'] == chain[i-1]['score']
    prompts = [r['messages'][1]['content'] for r in judge.requests if 'paragraph_scores' not in r['messages'][1]['content']]
    for text, entry in zip(prompts, chain):
        if entry['position'] > 1:
            assert f"Reference document ID: {entry['reference_doc_id']}" in text
            assert f"Reference Global Readability score: {entry['reference_score']:.4f}" in text
    summary = [json.loads(x) for x in (out / 'document_summary.jsonl').read_text().splitlines()]
    predicted = pd.read_csv(out / 'predictions/srad.tsv', sep='\t').set_index('answerId')
    for row in summary:
        assert len(row['section_repeat_scores']) == 3
        expected_global = round(sum(x['score'] for x in chain if x['doc_id'] == row['doc_id']) / 3, 4)
        assert row['global_score'] == expected_global
        expected = round(.3 * row['section_score'] + .7 * expected_global, 4)
        assert predicted.loc[row['doc_id'], 'readability_score'] == expected
    run(p, judge)
    assert len(judge.requests) == 18


def test_failed_question_is_not_reported_as_success(prepared):
    p, _, _, out = prepared
    class Broken(Judge):
        def create(self, **kwargs):
            raise ValueError('fixture failure')
    with pytest.raises(RuntimeError, match='SRAD failed'):
        run(p, Broken())
    assert not (out / 'predictions/srad.tsv').exists()


def test_position_weights_and_section_split(prepared):
    assert ev.readability_position_weights(1) == [1.0]
    assert ev.readability_position_weights(4) == [1.0, .8, .6, .4]
    assert ev.weighted_average([10, 0], [1, .4]) == 7.1429
    paras = ['Introduction', 'Explanation. ' * 60, 'References', 'https://example.com']
    blocks, _, dropped = sections.build_readability_sections(ev, paras)
    assert dropped == 2
    assert all('https://example.com' not in x['text'] for x in blocks)


def test_resume_rejects_changed_document(prepared):
    p, reports, rubrics, out = prepared
    ev._write_json_atomic(out / 'manifest.json', p['manifest'])
    (reports / 'Report1/Doc_001.docx').write_bytes(b'changed input')
    with pytest.raises(ValueError, match='changed'):
        configure(ROOT / 'configs/srad-deepseek-v4-flash-paper.yaml', reports, rubrics, out)


def test_context_fallback_is_recorded(prepared, monkeypatch):
    p, _, _, out = prepared
    original = ev.evaluate_global_readability_prompt
    def context_limited(**kwargs):
        if '[REFERENCE_DOCUMENT_START]' in kwargs['prompt']:
            raise ev.ContextLengthError('fixture context limit')
        return original(**kwargs)
    monkeypatch.setattr(ev, 'evaluate_global_readability_prompt', context_limited)
    run(p, Judge())
    chain = [json.loads(x) for x in (out / 'global_chain_scores.jsonl').read_text().splitlines()]
    assert sum(x['context_fallback'] for x in chain) == 6
    assert all(x['mode'] == 'independent_context_fallback' for x in chain if x['context_fallback'])
